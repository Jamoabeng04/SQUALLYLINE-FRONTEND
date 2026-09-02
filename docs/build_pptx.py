# -*- coding: utf-8 -*-
"""
Build the Squally Line video deck as a native PowerPoint (.pptx).

Mirrors docs/slides.html — 13 slides, 16:9, white/black/gold palette.
Uses fonts reliably present on Windows (Georgia / Segoe UI / Consolas) so the
file opens correctly in the team's PowerPoint without web-font downloads.

Run:  python docs/build_pptx.py
Out:  docs/Squally-Line-Slides.pptx
"""

import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

# ---- palette (matches slides.html light theme) ----------------------------
PAPER       = RGBColor(0xFB, 0xFA, 0xF7)
PANEL       = RGBColor(0xF4, 0xF1, 0xEA)
PANEL2      = RGBColor(0xFF, 0xFF, 0xFF)
INK         = RGBColor(0x17, 0x14, 0x0F)
INK_SOFT    = RGBColor(0x40, 0x3A, 0x2F)
MUTED       = RGBColor(0x6B, 0x65, 0x58)
FAINT       = RGBColor(0x9A, 0x94, 0x84)
LINE        = RGBColor(0xE1, 0xDA, 0xCA)
LINE_STRONG = RGBColor(0xD2, 0xC9, 0xB4)
GOLD        = RGBColor(0xD4, 0xAF, 0x37)
GOLD_INK    = RGBColor(0x9A, 0x7B, 0x2E)
GOLD_WASH   = RGBColor(0xF6, 0xEE, 0xD4)

SERIF = "Georgia"
SANS  = "Segoe UI"
MONO  = "Consolas"

# ---- geometry (inches) -----------------------------------------------------
SW, SH = 13.333, 7.5
MX = 0.92                      # left / right margin
W  = SW - 2 * MX               # usable width
EY_Y = 0.52                    # eyebrow baseline
H2_Y = 0.95                    # title
CONTENT_TOP = 2.15

prs = Presentation()
prs.slide_width  = Inches(SW)
prs.slide_height = Inches(SH)
BLANK = prs.slide_layouts[6]


# ---- helpers ---------------------------------------------------------------
def parse(s):
    """'plain **bold** plain' -> [(text, is_bold), ...]"""
    out = []
    for p in re.split(r'(\*\*.*?\*\*)', s):
        if not p:
            continue
        if p.startswith('**') and p.endswith('**'):
            out.append((p[2:-2], True))
        else:
            out.append((p, False))
    return out


def new_slide():
    s = prs.slides.add_slide(BLANK)
    s.background.fill.solid()
    s.background.fill.fore_color.rgb = PAPER
    return s


def add_text(slide, x, y, w, h, paras, align=PP_ALIGN.LEFT,
             anchor=MSO_ANCHOR.TOP, line_spacing=None, space_after=0, wrap=True):
    """paras = list of paragraphs; each paragraph = list of (text, spec) runs."""
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = wrap
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    for pi, para in enumerate(paras):
        p = tf.paragraphs[0] if pi == 0 else tf.add_paragraph()
        p.alignment = align
        p.space_before = Pt(0)
        p.space_after = Pt(space_after)
        if line_spacing:
            p.line_spacing = line_spacing
        for text, spec in para:
            r = p.add_run()
            r.text = text
            f = r.font
            f.name = spec.get('font', SANS)
            f.size = Pt(spec.get('size', 14))
            f.bold = spec.get('bold', False)
            f.italic = spec.get('italic', False)
            f.color.rgb = spec.get('color', INK)
    return tb


def rrect(slide, x, y, w, h, fill, line_color=None, line_w=1.0, radius=0.07):
    sp = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE,
                                Inches(x), Inches(y), Inches(w), Inches(h))
    sp.shadow.inherit = False
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line_color is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line_color
        sp.line.width = Pt(line_w)
    try:
        sp.adjustments[0] = radius
    except (IndexError, KeyError):
        pass
    return sp


def spec(size, color=INK_SOFT, bold=False, italic=False, font=SANS):
    return {'font': font, 'size': size, 'color': color, 'bold': bold, 'italic': italic}


def eyebrow(slide, label, num):
    add_text(slide, MX, EY_Y, W * 0.7, 0.3,
             [[(label.upper(), spec(11, GOLD_INK, True, font=MONO))]])
    add_text(slide, MX, EY_Y, W, 0.3,
             [[(num, spec(11, FAINT, font=MONO))]], align=PP_ALIGN.RIGHT)


def title(slide, text, size=40, color=INK, y=H2_Y, w=W, align=PP_ALIGN.LEFT):
    add_text(slide, MX, y, w, 1.15,
             [[(text, spec(size, color, font=SERIF))]],
             align=align, line_spacing=1.0)


def bullets(items, size=12):
    """items -> list of paragraphs, each with a gold dot + parsed runs."""
    paras = []
    for it in items:
        para = [('●  ', spec(size * 0.82, GOLD))]
        for text, b in parse(it):
            para.append((text, spec(size, INK if b else INK_SOFT, b)))
        paras.append(para)
    return paras


def column(slide, x, y, w, h, header, items, accent=False, size=12):
    rrect(slide, x, y, w, h,
          GOLD_WASH if accent else PANEL,
          line_color=GOLD if accent else LINE, line_w=1.0)
    pad = 0.3
    if header:
        add_text(slide, x + pad, y + pad, w - 2 * pad, 0.3,
                 [[(header.upper(), spec(10.5, MUTED, True, font=MONO))]])
    add_text(slide, x + pad, y + pad + (0.52 if header else 0),
             w - 2 * pad, h - 2 * pad - (0.52 if header else 0),
             bullets(items, size), line_spacing=1.12, space_after=7)


def two_col(eb, num, ttl, ha, ia, hb, ib, accent_a=False, accent_b=False,
            lead=None):
    s = new_slide()
    eyebrow(s, eb, num)
    title(s, ttl)
    bottom = 6.05 if lead else 6.9
    colH = bottom - CONTENT_TOP
    colW = (W - 0.42) / 2
    column(s, MX, CONTENT_TOP, colW, colH, ha, ia, accent_a)
    column(s, MX + colW + 0.42, CONTENT_TOP, colW, colH, hb, ib, accent_b)
    if lead:
        add_text(s, MX, bottom + 0.28, W, 0.6,
                 [[(t, spec(15, INK if b else INK_SOFT, b)) for t, b in parse(lead)]],
                 line_spacing=1.3)
    return s


def node(slide, x, y, w, h, t, sub, ext=False):
    rrect(slide, x, y, w, h,
          GOLD_WASH if ext else PANEL2,
          line_color=GOLD if ext else LINE_STRONG, line_w=1.25)
    add_text(slide, x, y + h * 0.20, w, h * 0.42,
             [[(t, spec(15, INK, True))]], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)
    add_text(slide, x, y + h * 0.56, w, h * 0.34,
             [[(sub, spec(9.5, MUTED, font=MONO))]], align=PP_ALIGN.CENTER,
             anchor=MSO_ANCHOR.MIDDLE)


def connector(slide, x, y, w, h, lbl):
    add_text(slide, x, y, w, h,
             [[('→', spec(20, GOLD_INK, font=MONO))],
              [(lbl, spec(8.5, MUTED, font=MONO))]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.0)


# ===========================================================================
# SLIDE 1 — COVER
# ===========================================================================
s = new_slide()
add_text(s, MX, 0.6, W, 0.3,
         [[('COE 454 · SE II', spec(11, FAINT, font=MONO))]],
         align=PP_ALIGN.RIGHT)
rrect(s, MX, 2.15, 1.7, 0.06, GOLD, radius=0)
add_text(s, MX - 0.03, 2.35, W, 2.9,
         [[('Squally', spec(96, INK, font=SERIF))],
          [('Line', spec(96, INK, font=SERIF)), ('.', spec(96, GOLD_INK, font=SERIF))]],
         line_spacing=0.92)
add_text(s, MX, 5.55, 8.6, 0.9,
         [[('A web platform built for a real fashion-design client — '
            'from requirements to a live, deployed system.', spec(17, MUTED))]],
         line_spacing=1.3)
rrect(s, MX, 6.55, 6.35, 0.55, PAPER, line_color=GOLD, line_w=1.25, radius=0.5)
add_text(s, MX, 6.55, 6.35, 0.55,
         [[('squallyline-frontend-production.up.railway.app', spec(13, INK, font=MONO))]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ===========================================================================
# SLIDE 2 — CLIENT & PROBLEM
# ===========================================================================
two_col(
    'The Client', '02 / 13', 'A studio run on paper and WhatsApp',
    'The client', [
        '**[Client name]**, fashion designer',
        'Based in **[City]**, Ghana',
        'Centres: **bespoke tailoring · ready-to-wear · consultations & fittings**',
    ],
    'The problem', [
        'Measurements kept in **paper notebooks**',
        'Orders & quotes negotiated over **WhatsApp**',
        "Production progress tracked **in the tailor's head**",
        "Customers **couldn't browse, save measurements, or follow an order**",
    ],
)

# ===========================================================================
# SLIDE 3 — SDLC
# ===========================================================================
s = new_slide()
eyebrow(s, 'The Process We Followed', '03 / 13')
title(s, 'A deliberate software development life cycle')
phases = [
    ('01', 'Problem definition', False),
    ('02', 'Requirements & analysis', True),
    ('03', 'Planning & design', False),
    ('04', 'Implementation', False),
    ('05', 'Testing', False),
    ('06', 'Deployment', False),
    ('07', 'Maintenance', False),
]
n = len(phases)
gap = 0.18
bw = (W - gap * (n - 1)) / n
by, bh = 2.35, 2.5
for i, (idx, nm, hi) in enumerate(phases):
    bx = MX + i * (bw + gap)
    rrect(s, bx, by, bw, bh, GOLD_WASH if hi else PANEL,
          line_color=GOLD if hi else LINE, line_w=1.0)
    add_text(s, bx, by + 0.24, bw, 0.3,
             [[(idx, spec(9, GOLD_INK, font=MONO))]], align=PP_ALIGN.CENTER)
    add_text(s, bx + 0.1, by + 0.3, bw - 0.2, bh - 0.5,
             [[(nm, spec(11.5, GOLD_INK if hi else INK, True))]],
             align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, line_spacing=1.05)
add_text(s, MX, by + bh + 0.32, W, 0.6,
         [[(t, spec(15, INK if b else INK_SOFT, b)) for t, b in
           parse('Run **Agile & incrementally**, prioritised with **MoSCoW**, '
                 'shipping a working MVP first.')]],
         line_spacing=1.3)

# ===========================================================================
# SLIDE 4 — REQUIREMENTS METHODS
# ===========================================================================
two_col(
    'Requirements Elicitation', '04 / 13', 'The methods we chose — and why',
    'Primary', [
        '**Stakeholder interviews** — semi-structured, open-ended sessions with the designer',
        '**Observation / contextual inquiry** — watched him measure a client and run an order end to end',
    ],
    'Supporting', [
        '**Document analysis** — his paper measurement books & WhatsApp order chats',
        '**Iterative prototyping** — mock-ups reviewed and refined with him',
    ],
    accent_a=True,
    lead='Chosen to fit **one expert client** — not a survey crowd.',
)

# ===========================================================================
# SLIDE 5 — FR vs NFR
# ===========================================================================
two_col(
    'Requirements', '05 / 13', 'Functional & non-functional',
    'Functional — what it does', [
        'Measurement profiles & history',
        'Catalogue: products & made-to-order styles',
        'Cart, checkout & online payment',
        'Quote → agreement → order workflow',
        'Production tracking & analytics',
    ],
    'Non-functional — how well', [
        '**Security** — money & personal data',
        '**Performance & scalability**',
        '**Availability & reliability**',
        '**Usability**',
        '**Compliance** — Ghana payments, cedis',
    ],
)

# ===========================================================================
# SLIDE 6 — ARCHITECTURE
# ===========================================================================
s = new_slide()
eyebrow(s, 'System Architecture', '06 / 13')
title(s, 'A three-tier client–server web app')
nw, nh = 2.30, 1.25
cw = 0.72
row_y = 2.65
nodes = [
    ('Browser', 'customer · staff'),
    ('React SPA', 'React 19 · Router · three.js'),
    ('Django + DRF', 'Gunicorn · WhiteNoise'),
    ('PostgreSQL', 'Railway'),
]
conns = ['HTTPS', 'REST · JWT', 'SQL']
x = MX
for i, (t, sub) in enumerate(nodes):
    node(s, x, row_y, nw, nh, t, sub)
    x += nw
    if i < len(conns):
        connector(s, x, row_y, cw, nh, conns[i])
        x += cw
# external systems row
ew = 2.8
tot = 2 * ew + cw
ex = MX + (W - tot) / 2
ext_y = 4.55
node(s, ex, ext_y, ew, 1.2, 'Paystack', 'payments · GHS · webhook', ext=True)
node(s, ex + ew + cw, ext_y, ew, 1.2, 'Railway', 'hosting · both tiers', ext=True)

# ===========================================================================
# SLIDE 7 — ADRs
# ===========================================================================
two_col(
    'Architecture Decision Records', '07 / 13', 'Decisions we recorded & can defend',
    None, [
        '**ADR-001** — Separate SPA + REST API, not a monolith',
        '**ADR-002** — Stateless JWT auth over server sessions',
        '**ADR-003** — Paystack gateway (Ghana market, cedis)',
    ],
    None, [
        '**ADR-004** — PostgreSQL in production',
        '**ADR-005** — Railway hosting for both tiers',
        '**ADR-006** — Permissive-licence-only dependencies (React MIT · Django BSD)',
    ],
    lead=None,
)

# ===========================================================================
# SLIDE 8 — ROLES
# ===========================================================================
s = new_slide()
eyebrow(s, 'Role-Based Access', '08 / 13')
title(s, 'Three roles, three different worlds')
roles = [
    ('Customer', 'shop', ['Browse & buy', 'Measurements & people',
                          'Requests & orders', 'Accept proposals'], False),
    ('Apprentice', 'shop floor', ['Review consultations', 'Advance production',
                                  'Manage orders', 'Day-to-day ops'], False),
    ('Admin', 'owner', ['Everything staff can', 'Catalogue & pricing',
                        'Compose proposals', 'Analytics & export'], True),
]
cgap = 0.4
cw3 = (W - 2 * cgap) / 3
cy, ch = CONTENT_TOP, 4.55
for i, (rh, rs, items, accent) in enumerate(roles):
    cx = MX + i * (cw3 + cgap)
    rrect(s, cx, cy, cw3, ch, GOLD_WASH if accent else PANEL,
          line_color=GOLD if accent else LINE, line_w=1.0)
    pad = 0.32
    add_text(s, cx + pad, cy + pad, cw3 - 2 * pad, 0.6,
             [[(rh, spec(23, INK, font=SERIF, bold=True))]])
    add_text(s, cx + pad, cy + pad + 0.62, cw3 - 2 * pad, 0.3,
             [[(rs.upper(), spec(9, GOLD_INK, True, font=MONO))]])
    add_text(s, cx + pad, cy + pad + 1.05, cw3 - 2 * pad, ch - pad - 1.1,
             bullets(items, 12.5), line_spacing=1.15, space_after=8)

# ===========================================================================
# SLIDE 9 — SECURITY / NFR
# ===========================================================================
two_col(
    'Non-Functional Requirements', '09 / 13', 'The qualities behind the features',
    'Security', [
        'JWT with rotating, blacklistable refresh tokens',
        '**Role permissions on the server**, not just hidden UI',
        'Strong password policy',
        'Paystack webhooks **signature-verified**',
    ],
    'Performance & reliability', [
        'Paginated lists · images optimised to **WebP**',
        'Silent token refresh',
        'Checkout & accept run in **atomic transactions**',
        '**Idempotent** webhook — no double-charge',
    ],
)

# ===========================================================================
# SLIDE 10 — LIMITATIONS
# ===========================================================================
two_col(
    'Honesty', '10 / 13', 'Limitations & problems we faced',
    'Not done yet', [
        'No email / SMS notifications (configured, not wired)',
        'Refunds modelled, no full workflow',
        'Some policy copy still static',
        'Top profile menu on placeholder data',
    ],
    'Challenges faced', [
        '**7 people, one codebase** — Git branches, merge conflicts',
        'Learning **three.js** for the 3D figure',
        '**Railway deploy** — env config, CORS, static files',
        '**Scope creep** — controlled with MoSCoW & a scope freeze',
    ],
)

# ===========================================================================
# SLIDE 11 — SCALING
# ===========================================================================
s = new_slide()
eyebrow(s, 'Scalability', '11 / 13')
title(s, 'What breaks first at 10,000 users?')
steps = [
    ('Harden', 'secrets & debug into env vars', True),
    ('Optimise', 'indexes, caching', False),
    ('Scale up', 'bigger instance', False),
    ('Scale out', 'many API nodes — JWT already stateless', False),
    ('Decompose', 'media to CDN, background worker', False),
]
ns = len(steps)
agap = 0.5
sw = (W - agap * (ns - 1)) / ns
sy, sh_ = 2.7, 1.9
x = MX
for i, (st, sd, first) in enumerate(steps):
    rrect(s, x, sy, sw, sh_, GOLD_WASH if first else PANEL,
          line_color=GOLD if first else LINE, line_w=1.0)
    add_text(s, x + 0.12, sy + 0.28, sw - 0.24, 0.5,
             [[(st, spec(15, INK, True))]], align=PP_ALIGN.CENTER)
    add_text(s, x + 0.14, sy + 0.82, sw - 0.28, sh_ - 0.95,
             [[(sd, spec(10.5, MUTED))]], align=PP_ALIGN.CENTER, line_spacing=1.2)
    x += sw
    if i < ns - 1:
        add_text(s, x, sy, agap, sh_,
                 [[('→', spec(22, GOLD_INK, font=MONO))]],
                 align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
        x += agap
add_text(s, MX, sy + sh_ + 0.34, W, 0.6,
         [[(t, spec(15, INK if b else INK_SOFT, b)) for t, b in
           parse('The signal to scale is **real demand, not ambition.**')]],
         line_spacing=1.3)

# ===========================================================================
# SLIDE 12 — ROADMAP & HANDOVER
# ===========================================================================
two_col(
    'What Comes Next', '12 / 13', 'Roadmap & handover',
    'Product roadmap', [
        'Real email · SMS · WhatsApp notifications',
        'Full refund flow',
        'Mobile app',
        'Deeper analytics',
    ],
    'Business & IP', [
        '7 developers = **joint authors** (Ghana Copyright Act 690)',
        '**IP assigned to client** on final payment',
        'Signed **MVP acceptance**',
        'Hosting & support **retainer**',
    ],
    accent_b=True,
)

# ===========================================================================
# SLIDE 13 — CLOSE
# ===========================================================================
s = new_slide()
add_text(s, MX, 1.6, W, 0.3,
         [[('THANK YOU', spec(11, GOLD_INK, True, font=MONO))]],
         align=PP_ALIGN.CENTER)
add_text(s, MX, 2.4, W, 1.9,
         [[('A real problem, solved with real', spec(46, INK, font=SERIF))],
          [('software engineering.', spec(46, INK, font=SERIF))]],
         align=PP_ALIGN.CENTER, line_spacing=1.0)
add_text(s, MX, 4.75, W, 0.5,
         [[(t, spec(16, MUTED, b)) for t, b in
           parse('Our thanks to **[Client name]** for trusting us with his studio.')]],
         align=PP_ALIGN.CENTER)
pill_w = 6.6
rrect(s, MX + (W - pill_w) / 2, 5.6, pill_w, 0.6, PAPER,
      line_color=GOLD, line_w=1.25, radius=0.5)
add_text(s, MX + (W - pill_w) / 2, 5.6, pill_w, 0.6,
         [[('squallyline-frontend-production.up.railway.app', spec(14, INK, font=MONO))]],
         align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)

# ---- save ------------------------------------------------------------------
out = os.path.join(os.path.dirname(os.path.abspath(__file__)), 'Squally-Line-Slides.pptx')
prs.save(out)
print('OK wrote %s (%d slides)' % (out, len(prs.slides._sldIdLst)))
